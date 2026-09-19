"""Prototype-only intake selection, persistence and export regressions."""

import json
from datetime import date, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from backend.app import create_app
from backend.extensions import db
from backend.models import EnrollmentDraft, User
from backend.models.finance import MasterInvoiceExport
from backend.services.course_catalogue import CATALOGUE_PATH, CourseCatalogue
from backend.services.master_invoice_service import HEADERS, SHEET_NAME, TABLE_NAME
from backend.tests.conftest import authenticated_client, send


DEMO_RECORDS = (
    ("demo-2026-10-19", "2026-10-19", "2026-10-20", "19–20 October 2026"),
    ("demo-2026-11-16", "2026-11-16", "2026-11-17", "16–17 November 2026"),
    ("demo-2026-12-14", "2026-12-14", "2026-12-15", "14–15 December 2026"),
)


def demo_app(tmp_path, fake_ai, *, catalogue=None, enabled=True, name="demo.db"):
    return create_app(
        {
            "TESTING": True,
            "SECRET_KEY": "demo-intake-test-secret",
            "OPENAI_API_KEY": "",
            "SQLALCHEMY_DATABASE_URI": f"sqlite:///{(tmp_path / name).as_posix()}",
            "ENABLE_DEMO_INTAKES": enabled,
        },
        ai_service=fake_ai,
        catalogue=catalogue,
    )


def begin_intake_selection(client):
    conversation_id = client.post("/api/conversations").get_json()["conversation_id"]
    send(client, conversation_id, "I want to enroll")
    reply = send(client, conversation_id, "Yes, this course").get_json()
    return conversation_id, reply


def test_demo_records_are_structured_unique_valid_and_consecutive():
    catalogue = CourseCatalogue(enable_demo_intakes=True)
    records = catalogue.demo_intakes

    assert [
        (item["id"], item["start_date"], item["end_date"], item["display_date"])
        for item in records
    ] == list(DEMO_RECORDS)
    assert len({item["id"] for item in records}) == 3
    assert all(item["is_demo"] is True for item in records)
    assert all(item["status"] == "open" for item in records)
    for item in records:
        assert date.fromisoformat(item["end_date"]) == (
            date.fromisoformat(item["start_date"]) + timedelta(days=1)
        )


def test_demo_dates_are_opt_in_and_hidden_by_default(tmp_path, fake_ai):
    disabled = CourseCatalogue()
    enabled = CourseCatalogue(enable_demo_intakes=True)
    assert disabled.intake_dates == []
    assert "demo-2026" not in disabled.agent_context()
    assert enabled.intake_dates == [record[1] for record in DEMO_RECORDS]

    app = demo_app(tmp_path, fake_ai, enabled=False)
    client = authenticated_client(app)
    created = client.post("/api/conversations").get_json()
    assert created["demo_intakes_enabled"] is False
    offer = send(client, created["conversation_id"], "I want to enroll").get_json()
    assert "Awaiting staff confirmation" in offer["message"]["content"]
    assert "19–20 October 2026" not in offer["message"]["content"]

    env_example = (CATALOGUE_PATH.parents[1] / ".env.example").read_text(
        encoding="utf-8"
    )
    assert "ENABLE_DEMO_INTAKES=false" in env_example


def test_enabled_flow_labels_and_lists_demo_intakes(tmp_path, fake_ai):
    app = demo_app(tmp_path, fake_ai)
    client = authenticated_client(app)
    conversation_id, selection = begin_intake_selection(client)

    assert selection["enrollment"]["demo_intakes_enabled"] is True
    content = selection["message"]["content"]
    assert "Demo intake dates — prototype only" in content
    assert "Please choose a demo intake" in content
    for index, record in enumerate(DEMO_RECORDS, start=1):
        assert f"{index}. {record[3]}" in content
    assert "Reply with 1, 2, 3, or the intake date" in content
    assert conversation_id


@pytest.mark.parametrize(
    ("selection", "expected_id"),
    [("1", DEMO_RECORDS[0][0]), ("2", DEMO_RECORDS[1][0]), ("3", DEMO_RECORDS[2][0])],
)
def test_numbered_selection_persists_structured_intake(
    tmp_path, fake_ai, selection, expected_id
):
    app = demo_app(tmp_path, fake_ai)
    client = authenticated_client(app)
    conversation_id, _ = begin_intake_selection(client)
    response = send(client, conversation_id, selection).get_json()

    selected = response["enrollment"]["selected_intake"]
    assert selected["id"] == expected_id
    assert selected["is_demo"] is True
    with app.app_context():
        draft = EnrollmentDraft.query.one()
        expected = next(record for record in DEMO_RECORDS if record[0] == expected_id)
        assert draft.intake_id == expected_id
        assert draft.intake_start_date.isoformat() == expected[1]
        assert draft.intake_end_date.isoformat() == expected[2]
        assert draft.preferred_intake_date == draft.intake_start_date
        assert draft.intake_is_demo is True


@pytest.mark.parametrize(
    ("selection", "expected_id"),
    [
        ("2026-10-19", DEMO_RECORDS[0][0]),
        ("19–20 October 2026", DEMO_RECORDS[0][0]),
        ("the November intake", DEMO_RECORDS[1][0]),
        ("I choose 16 November", DEMO_RECORDS[1][0]),
    ],
)
def test_date_range_and_natural_language_selection(
    tmp_path, fake_ai, selection, expected_id
):
    app = demo_app(tmp_path, fake_ai)
    client = authenticated_client(app)
    conversation_id, _ = begin_intake_selection(client)
    response = send(client, conversation_id, selection).get_json()
    assert response["enrollment"]["selected_intake"]["id"] == expected_id


def test_arbitrary_closed_and_ambiguous_intakes_are_rejected(tmp_path, fake_ai):
    base = json.loads(CATALOGUE_PATH.read_text(encoding="utf-8"))
    base["course"]["demo_intakes"][1]["status"] = "closed"
    base["course"]["demo_intakes"].append(
        {
            "id": "demo-2026-11-23",
            "start_date": "2026-11-23",
            "end_date": "2026-11-24",
            "display_date": "23–24 November 2026",
            "status": "open",
            "is_demo": True,
        }
    )
    path = tmp_path / "intake-edge-cases.json"
    path.write_text(json.dumps(base), encoding="utf-8")
    catalogue = CourseCatalogue(path, enable_demo_intakes=True)

    closed_app = demo_app(tmp_path, fake_ai, catalogue=catalogue, name="closed.db")
    closed_client = authenticated_client(closed_app)
    conversation_id, _ = begin_intake_selection(closed_client)
    closed = send(closed_client, conversation_id, "2026-11-16").get_json()
    assert "intake is closed" in closed["message"]["content"]
    assert closed["enrollment"]["selected_intake"] is None

    ambiguous = send(closed_client, conversation_id, "the November intake").get_json()
    assert "more than one intake" in ambiguous["message"]["content"]
    arbitrary = send(closed_client, conversation_id, "Intake: 22 October 2026").get_json()
    assert "not one of the available intakes" in arbitrary["message"]["content"]


def test_intake_survives_faq_greeting_reload_logout_and_is_owner_scoped(
    tmp_path, fake_ai
):
    app = demo_app(tmp_path, fake_ai)
    owner = authenticated_client(app, email="demo-owner@example.com")
    conversation_id, _ = begin_intake_selection(owner)
    send(owner, conversation_id, "2")

    faq = send(owner, conversation_id, "Can I use SkillsFuture?").get_json()
    assert "SkillsFuture" in faq["message"]["content"]
    greeting = send(owner, conversation_id, "Good morning").get_json()
    assert greeting["message"]["content"].startswith("Good morning")
    history = owner.get(f"/api/conversations/{conversation_id}/messages").get_json()
    assert history["enrollment"]["selected_intake"]["id"] == DEMO_RECORDS[1][0]

    assert owner.post("/api/auth/logout").status_code == 200
    login = owner.post(
        "/api/auth/login",
        json={"email": owner.test_email, "password": owner.test_password},
    )
    assert login.status_code == 200
    resumed = owner.get(f"/api/conversations/{conversation_id}/messages")
    assert resumed.get_json()["enrollment"]["selected_intake"]["id"] == DEMO_RECORDS[1][0]

    other = authenticated_client(app, email="other-demo-user@example.com")
    assert other.get(f"/api/conversations/{conversation_id}/messages").status_code == 404
    with app.app_context():
        draft = EnrollmentDraft.query.one()
        owner_user = User.query.filter_by(email="demo-owner@example.com").one()
        assert draft.conversation.user_id == owner_user.id


def test_intake_correction_changes_only_intake_fields(tmp_path, fake_ai):
    app = demo_app(tmp_path, fake_ai)
    client = authenticated_client(app)
    conversation_id, _ = begin_intake_selection(client)
    send(client, conversation_id, "1")
    send(client, conversation_id, "Full name: Test Student")

    corrected = send(client, conversation_id, "Correct intake: 3").get_json()
    assert corrected["enrollment"]["selected_intake"]["id"] == DEMO_RECORDS[2][0]
    with app.app_context():
        draft = EnrollmentDraft.query.one()
        assert draft.full_name == "Test Student"
        assert draft.nric is None


def _complete_paynow_enrollment(client, conversation_id):
    send(client, conversation_id, "1")
    send(
        client,
        conversation_id,
        "Full name: Test Student; NRIC/FIN: S1234567D; "
        "Date of birth: 15 May 2000; Email: test.student@example.com; "
        "Mobile number: 91234567",
    )
    send(client, conversation_id, "PayNow only")
    return send(client, conversation_id, "S$600").get_json()


def test_confirmation_and_excel_export_are_single_owner_scoped_and_idempotent(
    tmp_path, fake_ai
):
    app = demo_app(tmp_path, fake_ai)
    client = authenticated_client(app, email="confirmed-demo@example.com")
    conversation_id, _ = begin_intake_selection(client)
    review = _complete_paynow_enrollment(client, conversation_id)

    assert review["enrollment"]["status"] == "awaiting_confirmation"
    assert "Selected demo intake: 19–20 October 2026" in review["message"]["content"]
    assert "prototype testing" in review["message"]["content"]
    service = app.extensions["master_invoice"]
    assert not service.workbook_path.exists()
    with app.app_context():
        assert MasterInvoiceExport.query.count() == 0

    first = send(client, conversation_id, "Confirm").get_json()
    second = send(client, conversation_id, "Confirm").get_json()
    assert first["enrollment"]["status"] == "confirmed"
    assert second["enrollment"]["status"] == "confirmed"

    with app.app_context():
        draft = EnrollmentDraft.query.one()
        assert draft.conversation.user.email == "confirmed-demo@example.com"
        assert draft.confirmed_at is not None
        assert MasterInvoiceExport.query.count() == 1
        assert MasterInvoiceExport.query.one().attempts == 1

    workbook = load_workbook(service.workbook_path)
    sheet = workbook[SHEET_NAME]
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet.max_row == 2
    assert sheet[2][3].value == "2-Day Basic Certificate in Dental Assisting"
    assert sheet[2][4].value == "19–20 October 2026"
    assert sheet[2][6].value == 600
    assert sheet[2][7].value == "S*******D"
    assert sheet[2][15].value == "PayNow"
    assert sheet[2][16].value == 600
    assert sheet[2][21].value == draft.id
    assert sheet[2][22].value == DEMO_RECORDS[0][0]
    assert sheet[2][23].value == "Yes"
    assert sheet[2][24].value == "Exported"
    assert TABLE_NAME in sheet.tables
