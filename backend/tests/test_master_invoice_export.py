"""Secure Master Invoice List export and staff-access regressions."""

from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

from backend.app import create_app
from backend.extensions import db
from backend.models import EnrollmentDraft, utc_now
from backend.models.finance import MasterInvoiceAudit, MasterInvoiceExport
from backend.services.master_invoice_service import HEADERS, SHEET_NAME, TABLE_NAME
from backend.tests.conftest import authenticated_client


def export_app(tmp_path, fake_ai):
    return create_app(
        {
            "TESTING": True,
            "SECRET_KEY": "master-invoice-test-secret",
            "SQLALCHEMY_DATABASE_URI": f"sqlite:///{(tmp_path / 'export.db').as_posix()}",
            "MASTER_INVOICE_PATH": str(tmp_path / "private" / "Master_Invoice_List.xlsx"),
            "STAFF_ACCOUNTS": {
                "accountant": {
                    "password_hash": generate_password_hash("staff-password"),
                    "role": "accountant",
                },
                "operator": {
                    "password_hash": generate_password_hash("staff-password"),
                    "role": "operations",
                },
            },
        },
        ai_service=fake_ai,
    )


def add_draft(app, client, *, confirmed=True, name="Test Student", split=True):
    conversation_id = client.post("/api/conversations").get_json()["conversation_id"]
    with app.app_context():
        draft = EnrollmentDraft(
            conversation_id=conversation_id,
            course="2-Day Basic Certificate in Dental Assisting",
            course_fee=Decimal("600"),
            full_name=name,
            nric="S1234567D",
            date_of_birth=date(2000, 5, 15),
            email="test.student@example.com",
            mobile_number="91234567",
            payment_method="skillsfuture_paynow" if split else "paynow",
            skillsfuture_amount=Decimal("400") if split else Decimal("0"),
            paynow_amount=Decimal("200") if split else Decimal("600"),
            status="awaiting_course_date" if confirmed else "collecting",
            confirmed_at=utc_now() if confirmed else None,
        )
        db.session.add(draft)
        db.session.commit()
        return draft.id


def staff_login(client, username="accountant"):
    response = client.post(
        "/api/staff/login",
        json={"username": username, "password": "staff-password"},
    )
    assert response.status_code == 200
    return response.get_json()["csrf_token"]


def csrf_post(client, path, token):
    return client.post(path, json={}, headers={"X-CSRF-Token": token})


def test_confirmed_rows_are_idempotent_formatted_and_rebuilt(tmp_path, fake_ai):
    app = export_app(tmp_path, fake_ai)
    participant = authenticated_client(app)
    first_id = add_draft(app, participant, name="=Formula Attempt")
    second_id = add_draft(app, participant, name="Unicode Student 李", split=False)
    add_draft(app, participant, confirmed=False, name="Unconfirmed Student")
    service = app.extensions["master_invoice"]

    with app.app_context():
        first = db.session.get(EnrollmentDraft, first_id)
        assert service.queue_and_export(first) is True
        assert service.queue_and_export(first) is True
        service.regenerate("accountant")
        assert MasterInvoiceExport.query.count() == 2

    workbook = load_workbook(service.workbook_path)
    sheet = workbook[SHEET_NAME]
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet.max_row == 3
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:Z3"
    assert TABLE_NAME in sheet.tables
    assert {sheet[2][21].value, sheet[3][21].value} == {first_id, second_id}
    assert sheet[2][5].value.startswith("'")
    assert sheet[2][6].number_format == '"S$"#,##0.00'
    assert sheet[2][13].number_format == "dd-mmm-yyyy"
    assert sheet[2][7].value == "S*******D"
    assert sheet[2][15].value == "SkillsFuture Credit"
    assert sheet[2][16].value == 400
    assert sheet[2][17].value == "PayNow"
    assert sheet[2][18].value == 200
    assert sheet[2][1].value is None
    assert sheet[2][2].value is None
    assert sheet[2][9].value is None

    service.workbook_path.unlink()
    with app.app_context():
        result = service.regenerate("accountant")
        assert result["confirmed_records"] == 2
        assert service.workbook_path.is_file()


def test_export_failure_keeps_confirmation_pending_and_retryable(
    tmp_path, fake_ai, monkeypatch
):
    app = export_app(tmp_path, fake_ai)
    participant = authenticated_client(app)
    enrollment_id = add_draft(app, participant)
    service = app.extensions["master_invoice"]
    original_save = service._atomic_save
    monkeypatch.setattr(
        service,
        "_atomic_save",
        lambda workbook: (_ for _ in ()).throw(PermissionError("locked")),
    )

    with app.app_context():
        draft = db.session.get(EnrollmentDraft, enrollment_id)
        assert service.queue_and_export(draft) is False
        db.session.expire_all()
        assert db.session.get(EnrollmentDraft, enrollment_id).confirmed_at is not None
        export = MasterInvoiceExport.query.filter_by(
            enrollment_id=enrollment_id
        ).one()
        assert export.status == "failed"
        assert export.last_error and "Retry" in export.last_error

    monkeypatch.setattr(service, "_atomic_save", original_save)
    with app.app_context():
        assert service.retry_failed("accountant") is True
        assert MasterInvoiceExport.query.one().status == "exported"


def test_master_invoice_endpoints_are_accounting_only_and_audited(
    tmp_path, fake_ai
):
    app = export_app(tmp_path, fake_ai)
    participant = authenticated_client(app)
    enrollment_id = add_draft(app, participant)
    with app.app_context():
        app.extensions["master_invoice"].queue_and_export(
            db.session.get(EnrollmentDraft, enrollment_id)
        )

    assert participant.get("/api/staff/master-invoice/download").status_code == 401
    operator = app.test_client()
    staff_login(operator, "operator")
    assert operator.get("/api/staff/master-invoice/download").status_code == 403

    accountant = app.test_client()
    token = staff_login(accountant)
    status = accountant.get("/api/staff/master-invoice/status")
    assert status.status_code == 200
    assert status.get_json()["master_invoice"]["confirmed_records"] == 1
    assert status.get_json()["master_invoice"]["pending_exports"] == 0
    assert status.get_json()["master_invoice"]["failed_exports"] == 0
    download = accountant.get("/api/staff/master-invoice/download")
    assert download.status_code == 200
    assert download.data.startswith(b"PK")
    assert "Master_Invoice_List.xlsx" in download.headers["Content-Disposition"]
    assert csrf_post(
        accountant, "/api/staff/master-invoice/regenerate", token
    ).status_code == 200
    with app.app_context():
        actions = [row.action for row in MasterInvoiceAudit.query.all()]
        assert "download" in actions
        assert "regenerate" in actions
