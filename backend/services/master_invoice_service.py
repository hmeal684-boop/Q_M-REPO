"""Private, recoverable Excel export of confirmed enrollment records."""

import os
import re
import tempfile
from io import BytesIO
from pathlib import Path

from filelock import FileLock, Timeout
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from backend.extensions import db
from backend.models import EnrollmentDraft, utc_now
from backend.models.finance import Invoice, MasterInvoiceAudit, MasterInvoiceExport
from backend.services.privacy import mask_nric


SHEET_NAME = "Master Invoice List"
WORKBOOK_NAME = "Master_Invoice_List.xlsx"
TABLE_NAME = "MasterInvoiceTable"
HEADERS = (
    "No.",
    "Invoice No.",
    "Course Code",
    "Course Name",
    "Lecturer/Course Date",
    "Student Name",
    "Amount Invoiced",
    "NRIC/FIN",
    "Invoice Date",
    "UCR No.",
    "Email",
    "Invoice Remark",
    "Handphone No.",
    "Date of Birth",
    "Private Note to Accounts",
    "Payment Mode 1",
    "Payment Amount 1",
    "Payment Mode 2",
    "Payment Amount 2",
    "Payment Mode 3",
    "Payment Amount 3",
    "Enrollment ID",
    "Intake ID",
    "Demo Intake",
    "Export Status",
    "Created At",
)


class MasterInvoiceService:
    def __init__(
        self, workbook_path, template_path="", lock_timeout=10, catalogue=None
    ):
        self.workbook_path = Path(workbook_path)
        self.template_path = Path(template_path) if template_path else None
        self.lock_timeout = float(lock_timeout)
        self.catalogue = catalogue

    def queue_and_export(self, enrollment, actor="system"):
        """Persist pending state first; never undo the confirmed enrollment."""
        if enrollment.confirmed_at is None:
            return False
        record = MasterInvoiceExport.query.filter_by(
            enrollment_id=enrollment.id
        ).one_or_none()
        if record is None:
            record = MasterInvoiceExport(
                enrollment_id=enrollment.id, status="pending"
            )
            db.session.add(record)
        elif record.status == "exported" and self.workbook_path.exists():
            return True
        else:
            record.status = "pending"
            record.last_error = None
        db.session.commit()
        return self._safe_rebuild(actor=actor, action="automatic")

    def regenerate(self, actor):
        return self._rebuild(actor=actor, action="regenerate")

    def retry_failed(self, actor):
        records = db.session.execute(
            db.select(MasterInvoiceExport).where(
                MasterInvoiceExport.status.in_(("pending", "failed"))
            )
        ).scalars()
        for record in records:
            record.status = "pending"
            record.last_error = None
        db.session.commit()
        return self._safe_rebuild(actor=actor, action="retry")

    def status(self):
        confirmed = self._confirmed_query().count()
        pending = MasterInvoiceExport.query.filter_by(status="pending").count()
        failed = MasterInvoiceExport.query.filter_by(status="failed").count()
        latest = db.session.execute(
            db.select(MasterInvoiceExport.exported_at)
            .where(MasterInvoiceExport.exported_at.is_not(None))
            .order_by(MasterInvoiceExport.exported_at.desc())
            .limit(1)
        ).scalar_one_or_none()
        return {
            "filename": WORKBOOK_NAME,
            "last_generated_at": _iso(latest),
            "confirmed_records": confirmed,
            "pending_exports": pending,
            "failed_exports": failed,
            "pending_or_failed": pending + failed,
            "download_available": self.workbook_path.is_file(),
        }

    def download(self, actor):
        if not self.workbook_path.is_file():
            raise LookupError("The Master Invoice List has not been generated yet")
        try:
            with FileLock(
                str(self.workbook_path) + ".lock", timeout=self.lock_timeout
            ):
                content = self.workbook_path.read_bytes()
        except Timeout as error:
            raise ValueError(
                "The workbook is currently in use. Please try again shortly."
            ) from error
        db.session.add(
            MasterInvoiceAudit(
                action="download",
                actor=actor,
                record_count=self._confirmed_query().count(),
            )
        )
        db.session.commit()
        return BytesIO(content)

    def _safe_rebuild(self, actor, action):
        try:
            self._rebuild(actor=actor, action=action)
            return True
        except Exception:
            db.session.rollback()
            records = db.session.execute(
                db.select(MasterInvoiceExport).where(
                    MasterInvoiceExport.status == "pending"
                )
            ).scalars()
            for record in records:
                record.status = "failed"
                record.attempts = (record.attempts or 0) + 1
                record.last_error = (
                    "Workbook generation failed or the file is in use. Retry is available."
                )
            db.session.add(
                MasterInvoiceAudit(
                    action=f"{action}_failed",
                    actor=actor,
                    record_count=self._confirmed_query().count(),
                )
            )
            db.session.commit()
            return False

    def _rebuild(self, actor, action):
        enrollments = self._confirmed_query().all()
        self._ensure_tracking(enrollments)
        self.workbook_path.parent.mkdir(parents=True, exist_ok=True)
        lock = FileLock(
            str(self.workbook_path) + ".lock", timeout=self.lock_timeout
        )
        try:
            with lock:
                workbook = self._new_workbook()
                worksheet = self._prepare_sheet(workbook)
                invoices = self._invoices(enrollments)
                for number, enrollment in enumerate(enrollments, start=1):
                    self._write_row(
                        worksheet,
                        number + 1,
                        number,
                        enrollment,
                        invoices.get(enrollment.id),
                    )
                self._format_sheet(worksheet, len(enrollments))
                self._atomic_save(workbook)
        except Timeout as error:
            raise ValueError(
                "The workbook is currently in use. The export remains pending."
            ) from error

        exported_at = utc_now()
        for enrollment in enrollments:
            record = MasterInvoiceExport.query.filter_by(
                enrollment_id=enrollment.id
            ).one()
            record.status = "exported"
            record.attempts = (record.attempts or 0) + 1
            record.last_error = None
            record.exported_at = exported_at
        db.session.add(
            MasterInvoiceAudit(
                action=action, actor=actor, record_count=len(enrollments)
            )
        )
        db.session.commit()
        return self.status()

    def _confirmed_query(self):
        return EnrollmentDraft.query.filter(
            EnrollmentDraft.confirmed_at.is_not(None),
            EnrollmentDraft.status != "cancelled",
        ).order_by(EnrollmentDraft.confirmed_at.asc(), EnrollmentDraft.id.asc())

    @staticmethod
    def _ensure_tracking(enrollments):
        known = {
            value
            for (value,) in db.session.execute(
                db.select(MasterInvoiceExport.enrollment_id)
            ).all()
        }
        for enrollment in enrollments:
            if enrollment.id not in known:
                db.session.add(
                    MasterInvoiceExport(
                        enrollment_id=enrollment.id, status="pending"
                    )
                )
        db.session.commit()

    @staticmethod
    def _invoices(enrollments):
        ids = [enrollment.id for enrollment in enrollments]
        if not ids:
            return {}
        rows = db.session.execute(
            db.select(Invoice).where(Invoice.enrollment_id.in_(ids))
        ).scalars()
        return {row.enrollment_id: row for row in rows}

    def _new_workbook(self):
        if self.template_path and self.template_path.is_file():
            return load_workbook(self.template_path)
        workbook = Workbook()
        workbook.active.title = SHEET_NAME
        return workbook

    @staticmethod
    def _prepare_sheet(workbook):
        if SHEET_NAME in workbook.sheetnames:
            worksheet = workbook[SHEET_NAME]
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    if cell.data_type != "f":
                        cell.value = None
        else:
            worksheet = workbook.create_sheet(SHEET_NAME)
        for column, header in enumerate(HEADERS, start=1):
            worksheet.cell(row=1, column=column, value=header)
        return worksheet

    def _write_row(self, worksheet, row_number, number, enrollment, invoice):
        modes = self._payment_modes(enrollment)
        amount_invoiced = (
            invoice.course_fee
            if invoice is not None
            else enrollment.course_fee
        )
        values = [
            number,
            invoice.number if invoice else None,
            self._safe_text(self._course_code(enrollment)),
            self._safe_text(enrollment.course),
            self._safe_text(self._intake_display(enrollment)),
            self._safe_text(enrollment.full_name),
            amount_invoiced,
            self._safe_text(mask_nric(enrollment.nric)),
            invoice.created_at.date() if invoice else None,
            None,
            self._safe_text(enrollment.email),
            None,
            self._safe_text(enrollment.mobile_number),
            enrollment.date_of_birth,
            None,
            self._safe_text(modes[0][0]) if len(modes) > 0 else None,
            modes[0][1] if len(modes) > 0 else None,
            self._safe_text(modes[1][0]) if len(modes) > 1 else None,
            modes[1][1] if len(modes) > 1 else None,
            self._safe_text(modes[2][0]) if len(modes) > 2 else None,
            modes[2][1] if len(modes) > 2 else None,
            enrollment.id,
            self._safe_text(enrollment.intake_id),
            "Yes" if enrollment.intake_is_demo else (
                "No" if enrollment.intake_id else None
            ),
            "Exported",
            enrollment.created_at,
        ]
        for column, value in enumerate(values, start=1):
            worksheet.cell(row=row_number, column=column, value=value)

    def _course_code(self, enrollment):
        if (
            self.catalogue is not None
            and enrollment.course == self.catalogue.course_name
        ):
            return self.catalogue.course.get("code")
        return None

    @staticmethod
    def _intake_display(enrollment):
        start = enrollment.intake_start_date or enrollment.preferred_intake_date
        end = enrollment.intake_end_date or start
        if start is None:
            return None
        if start == end:
            return f"{start.day} {start.strftime('%B %Y')}"
        if start.year == end.year and start.month == end.month:
            return f"{start.day}–{end.day} {start.strftime('%B %Y')}"
        if start.year == end.year:
            return f"{start.day} {start.strftime('%B')}–{end.day} {end.strftime('%B %Y')}"
        return (
            f"{start.day} {start.strftime('%B %Y')}–"
            f"{end.day} {end.strftime('%B %Y')}"
        )

    @staticmethod
    def _payment_modes(enrollment):
        modes = []
        if enrollment.skillsfuture_amount and enrollment.skillsfuture_amount > 0:
            modes.append(("SkillsFuture Credit", enrollment.skillsfuture_amount))
        if enrollment.paynow_amount and enrollment.paynow_amount > 0:
            modes.append(("PayNow", enrollment.paynow_amount))
        if enrollment.payment_method == "utap":
            modes.append(("UTAP reimbursement", None))
        return modes[:3]

    @staticmethod
    def _safe_text(value):
        if value is None:
            return None
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))
        if text.startswith(("=", "+", "-", "@")):
            return "'" + text
        return text

    @staticmethod
    def _format_sheet(worksheet, data_rows):
        worksheet.freeze_panes = "A2"
        final_row = data_rows + 1
        table_ref = f"A1:{get_column_letter(len(HEADERS))}{final_row}"
        worksheet.auto_filter.ref = table_ref
        widths = {
            "A": 7, "B": 18, "C": 15, "D": 38, "E": 22, "F": 28,
            "G": 18, "H": 16, "I": 16, "J": 15, "K": 32, "L": 25,
            "M": 18, "N": 16, "O": 28, "P": 22, "Q": 18, "R": 22,
            "S": 18, "T": 22, "U": 18, "V": 38, "W": 24, "X": 14,
            "Y": 16, "Z": 22,
        }
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="548235")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column, width in widths.items():
            worksheet.column_dimensions[column].width = width
        for row in range(2, final_row + 1):
            for column in (9, 14):
                worksheet.cell(row=row, column=column).number_format = "dd-mmm-yyyy"
            worksheet.cell(row=row, column=26).number_format = "dd-mmm-yyyy hh:mm"
            for column in (7, 17, 19, 21):
                worksheet.cell(row=row, column=column).number_format = '"S$"#,##0.00'
            for column in (12, 15):
                worksheet.cell(row=row, column=column).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )
        if TABLE_NAME in worksheet.tables:
            del worksheet.tables[TABLE_NAME]
        if data_rows:
            table = Table(displayName=TABLE_NAME, ref=table_ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium4",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            worksheet.add_table(table)

    def _atomic_save(self, workbook):
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix="master-invoice-",
                suffix=".xlsx",
                dir=self.workbook_path.parent,
                delete=False,
            ) as handle:
                temporary = Path(handle.name)
            workbook.save(temporary)
            os.replace(temporary, self.workbook_path)
        finally:
            if temporary and temporary.exists():
                temporary.unlink(missing_ok=True)


def _iso(value):
    return value.isoformat().replace("+00:00", "Z") if value else None


__all__ = [
    "HEADERS",
    "MasterInvoiceService",
    "SHEET_NAME",
    "TABLE_NAME",
    "WORKBOOK_NAME",
]
